import collections.abc
import pptx
from pptx.util import Pt, Inches
import os

prs = pptx.Presentation('ScholarSetu_Presentation.pptx')

# 1. Update text on Frontend Implementation (Index 6)
slide_fe = prs.slides[6]
slide_fe.shapes.title.text = "Frontend & Backend Implementation"
for shape in slide_fe.shapes:
    if shape.has_text_frame and shape != slide_fe.shapes.title:
        shape.text = (
            "Frontend (React/Vite): Component-based UI with responsive chat, real-time Hindi localization, and secure state management hooks.\n"
            "Backend (FastAPI): Separated routes handling deterministic DAG state machines for conversational flow.\n"
            "OCR & Matching Engine: Pillow+pytesseract for extraction; strict deterministic evaluator for eligibility.\n"
            "Resilience: Lazy Evaluation via context maps avoids eager evaluation crashes on mismatched types."
        )

# 2. Update text on Challenges (Index 10)
slide_chal = prs.slides[10]
slide_chal.shapes.title.text = "Challenges & Future Scope"
for shape in slide_chal.shapes:
    if shape.has_text_frame and shape != slide_chal.shapes.title:
        shape.text = (
            "Challenges Solved: Handled IPv6/IPv4 Vite proxy conflicts, state machine crashes via lazy evaluation, and strict Pydantic UI prop mapping.\n"
            "Future - DigiLocker & Aadhaar: True zero-upload document verification and biometric/OTP demographic mapping.\n"
            "Future - PFMS/DBT: Tracking direct benefit transfers directly in the dashboard.\n"
            "Future - Scale & Analytics: Federated deployment with predictive analytics for scheme effectiveness and dropout prevention."
        )

# 3. Delete slides 11, 9, 7
sldIdLst = prs.slides._sldIdLst
for idx in [11, 9, 7]:
    slide = prs.slides[idx]
    rId = slide.slide_id
    for sldId in sldIdLst:
        if sldId.id == rId:
            sldIdLst.remove(sldId)
            break

# The indices have now shifted to exact 10 slides:
# 5 is System Architecture
# 7 is End-to-End Workflow
target_slides = {
    5: 'system_architecture.png',
    7: 'workflow.png'
}

for i, slide in enumerate(prs.slides):
    if i in target_slides:
        img_path = target_slides[i]
        
        shapes_to_delete = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_text_frame", False) or getattr(shape, "shape_type", None) == 13:
                shapes_to_delete.append(shape)
                
        for shape in shapes_to_delete:
            old_elem = shape.element
            old_elem.getparent().remove(old_elem)
            
        if os.path.exists(img_path):
            from PIL import Image
            with Image.open(img_path) as im:
                img_w, img_h = im.size
                ratio = img_w / img_h
                
            max_w_inches = 9.0
            max_h_inches = 5.0
            
            if max_w_inches / max_h_inches > ratio:
                final_h = max_h_inches
                final_w = max_h_inches * ratio
            else:
                final_w = max_w_inches
                final_h = max_w_inches / ratio
                
            slide_w_inches = 10.0
            left = Inches((slide_w_inches - final_w) / 2)
            top = Inches(1.8)
            
            slide.shapes.add_picture(img_path, left, top, width=Inches(final_w), height=Inches(final_h))

    else:
        for shape in slide.shapes:
            if shape.has_text_frame:
                is_title = (shape == slide.shapes.title)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if is_title:
                            run.font.size = Pt(36)
                        else:
                            run.font.size = Pt(20)

prs.save('ScholarSetu_Presentation.pptx')
print("Presentation edited, formatted, and saved.")
