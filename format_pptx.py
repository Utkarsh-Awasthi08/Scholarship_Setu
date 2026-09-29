import collections.abc
import pptx
from pptx.util import Pt, Inches
import os

prs = pptx.Presentation('ScholarSetu_Presentation.pptx')

# We want to replace text in slides 6 (index 5) and 9 (index 8) with images
target_slides = {
    5: 'system_architecture.png',
    8: 'workflow.png'
}

for i, slide in enumerate(prs.slides):
    # If it is one of the image replacement slides, find the body shape, delete it, and insert image
    if i in target_slides:
        img_path = target_slides[i]
        
        # Keep title, remove other text shapes and pictures
        shapes_to_delete = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_text_frame", False) or getattr(shape, "shape_type", None) == 13:
                shapes_to_delete.append(shape)
                
        for shape in shapes_to_delete:
            old_elem = shape.element
            old_elem.getparent().remove(old_elem)
            
        if os.path.exists(img_path):
            from PIL import Image
            with Image.open(img_path) as im:
                img_w, img_h = im.size
                ratio = img_w / img_h
                
            max_w_inches = 9.0
            max_h_inches = 5.0
            
            # Decide whether to constrain by width or height
            if max_w_inches / max_h_inches > ratio:
                # Constrain by height
                final_h = max_h_inches
                final_w = max_h_inches * ratio
            else:
                # Constrain by width
                final_w = max_w_inches
                final_h = max_w_inches / ratio
                
            slide_w_inches = 10.0 # Standard 4:3 is 10x7.5, 16:9 is 13.33x7.5. Assuming 10 for centering
            left = Inches((slide_w_inches - final_w) / 2)
            top = Inches(1.8) # Push it down a bit so title is clear
            
            slide.shapes.add_picture(img_path, left, top, width=Inches(final_w), height=Inches(final_h))

    else:
        # Format font size for all text
        for shape in slide.shapes:
            if shape.has_text_frame:
                is_title = (shape == slide.shapes.title)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if is_title:
                            run.font.size = Pt(36)
                        else:
                            run.font.size = Pt(20)

prs.save('ScholarSetu_Presentation.pptx')
print("Presentation formatted and saved.")
