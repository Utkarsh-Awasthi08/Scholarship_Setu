import collections.abc
import pptx

prs = pptx.Presentation('ScholarSetu_Presentation.pptx')
for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            print(f"Text: {shape.text}")
